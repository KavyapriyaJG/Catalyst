"""Shared document processing utilities for chunking, embedding, and context retrieval."""
import os
import sys
import io
from pathlib import Path
from typing import Any, List, Dict
from langchain_core.documents import Document
import PyPDF2
from docx import Document as DocxDocument
from config import get_settings

def chunk_text(text: str, chunk_size: int | None = None, overlap: int | None = None) -> List[str]:
    """Split text into overlapping chunks for embedding."""
    s = get_settings()
    chunk_size = chunk_size if chunk_size is not None else s.CHUNK_SIZE
    overlap = overlap if overlap is not None else s.CHUNK_OVERLAP
    cleaned = text.strip()
    if not cleaned:
        return []
    chunks: List[str] = []
    start = 0
    step = max(1, chunk_size - overlap)
    while start < len(cleaned):
        end = min(len(cleaned), start + chunk_size)
        chunks.append(cleaned[start:end])
        start += step
    return chunks

def build_documents(supporting_documents: List[Dict[str, str]]) -> List[Document]:
    """Convert raw documents to LangChain Document objects with metadata."""
    documents: List[Document] = []
    for source_index, item in enumerate(supporting_documents, start=1):
        filename = item.get("filename", f"document-{source_index}.txt")
        content = item.get("content", "")
        for chunk_index, chunk in enumerate(chunk_text(content), start=1):
            documents.append(
                Document(
                    page_content=chunk,
                    metadata={"source": filename, "chunk": chunk_index},
                )
            )
    return documents

def get_embeddings() -> Any:
    """Load or initialize HuggingFace embeddings model. Caches locally."""
    try:
        from langchain_huggingface import HuggingFaceEmbeddings
        from sentence_transformers import SentenceTransformer
    except ImportError as error:
        raise RuntimeError(
            "Missing dependency 'langchain-huggingface'. Install with: "
            "pip install langchain-huggingface sentence-transformers"
        ) from error

    s = get_settings()
    model_name = s.HUGGINGFACE_EMBEDDING_MODEL
    local_model_path = str(s.MODELS_EMBEDDINGS_DIR)
    os.makedirs(local_model_path, exist_ok=True)

    if not os.listdir(local_model_path):
        model = SentenceTransformer(model_name)
        model.save(local_model_path)

    return HuggingFaceEmbeddings(model_name=local_model_path)


def extract_document_content(filepath: Path) -> str:
    """
    Extract text content from various document formats.
    
    Supports: PDF (all pages), DOCX, TXT, and other text formats.
    Used centrally by all agents (PRD, backlog, modernization) for document context.
    
    Args:
        filepath: Path to the document file
    
    Returns:
        Full extracted text content
    
    Raises:
        ValueError: If file cannot be read or format is not supported
    """
    if not filepath.exists():
        raise ValueError(f"File not found: {filepath}")
    
    try:
        suffix = filepath.suffix.lower()
        text = None
        
        if suffix == ".pdf":
            text_content = []
            with open(filepath, "rb") as pdf_file:
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                for page in pdf_reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text_content.append(extracted)
            text = "\n".join(text_content)
        
        elif suffix == ".docx":
            doc = DocxDocument(filepath)
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        
        else:
            # Fallback for TXT and other text formats
            try:
                text = filepath.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                text = filepath.read_text(encoding="latin-1", errors="ignore")
        
        if not text:
            raise ValueError(f"No text content found in {filepath.name}")
        
        return text
    
    except Exception as e:
        if not isinstance(e, ValueError):
            raise ValueError(f"Failed to extract content from {filepath.name}: {str(e)}")
        raise


def retrieve_context(query: str, documents: List[Document], top_k: int = 6) -> str:
    """Retrieve relevant document chunks using semantic similarity search."""
    if not documents:
        return ""
    try:
        from langchain_core.vectorstores import InMemoryVectorStore
    except ImportError as error:
        raise RuntimeError(
            "Missing dependency 'langchain-core'. Install with: "
            "pip install langchain-core"
        ) from error
    
    old_stdout = sys.stdout
    old_stderr = sys.stderr
    
    # Create StringIO objects with isatty method to prevent errors
    null_out = io.StringIO()
    null_out.isatty = lambda: False
    null_err = io.StringIO()
    null_err.isatty = lambda: False
    
    try:
        sys.stdout = null_out
        sys.stderr = null_err
        
        embeddings = get_embeddings()
        vectorstore = InMemoryVectorStore.from_documents(documents, embeddings)
        results = vectorstore.similarity_search(query, k=top_k)
    finally:
        sys.stdout = old_stdout
        sys.stderr = old_stderr
    
    return "\n".join([doc.page_content for doc in results])


def extract_documents_from_uploads(uploads_dir: Path, documents: list[dict]) -> list[dict[str, str]]:
    """Extract content from uploaded documents in the uploads directory."""
    parsed_documents: list[dict[str, str]] = []
    
    for doc in documents:
        doc_id = doc.get("id")
        doc_name = doc.get("name")
        if doc_id and doc_name and uploads_dir.exists():
            found = False
            for filepath in uploads_dir.iterdir():
                if filepath.is_file() and filepath.name.startswith(f"{doc_id}###"):
                    try:
                        content = extract_document_content(filepath)
                        parsed_documents.append({
                            "filename": doc_name,
                            "content": content,
                        })
                        found = True
                    except Exception as e:
                        raise ValueError(f"Failed to extract {doc_name}: {e}")
                    break
    
    return parsed_documents