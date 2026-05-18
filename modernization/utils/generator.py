import io

from docx import Document
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from modernization.models import ModernizationDocRecord


def generate_docx_export(
    doc: ModernizationDocRecord,
    include_generated_sections: bool = True,
) -> bytes:
    document = Document()

    title = document.add_heading(doc.name, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_format = title.runs[0]
    title_format.font.color.rgb = RGBColor(0, 102, 204)

    metadata_table = document.add_table(rows=5, cols=2)
    metadata_table.style = 'Light Grid Accent 1'

    metadata_rows = [
        ("Document ID", doc.id),
        ("Status", doc.status.upper()),
        ("Created", doc.created_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("Last Updated", doc.updated_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("Submitted By", doc.submitted_by or "Not submitted"),
    ]

    for i, (key, value) in enumerate(metadata_rows):
        cells = metadata_table.rows[i].cells
        cells[0].text = key
        cells[1].text = value

    document.add_paragraph()

    if doc.description:
        document.add_heading("Overview", 1)
        document.add_paragraph(doc.description)
        document.add_paragraph()

    document.add_heading("Modernization Strategy", 1)
    goals_para = document.add_paragraph(doc.modernization_goals)
    goals_para.style = 'Normal'

    document.add_paragraph()

    if doc.linked_prds:
        document.add_heading("Linked Requirements", 1)
        for prd_id in doc.linked_prds:
            document.add_paragraph(f"• {prd_id}", style='List Bullet')
        document.add_paragraph()

    if doc.source_assets:
        document.add_heading("Supporting Documentation", 1)
        for asset in doc.source_assets:
            filename = asset.get("filename", "unknown")
            size = asset.get("size", 0)
            size_mb = f"{size / (1024*1024):.2f} MB"
            document.add_paragraph(f"• {filename} ({size_mb})", style='List Bullet')
        document.add_paragraph()

    if include_generated_sections and doc.generated_sections:
        document.add_page_break()
        document.add_heading("Auto-Generated Analysis", 1)

        sections = doc.generated_sections

        if "executive_summary" in sections:
            summary = sections["executive_summary"]
            document.add_heading("Executive Summary", 2)
            document.add_paragraph(summary)
            document.add_paragraph()

        if "current_target_state" in sections:
            state = sections["current_target_state"]
            document.add_heading("Current State → Target State", 2)
            document.add_paragraph(state)
            document.add_paragraph()

        if "implementation_approach" in sections:
            approach = sections["implementation_approach"]
            document.add_heading("Implementation Approach", 2)
            document.add_paragraph(approach)
            document.add_paragraph()

        if "risks_mitigation" in sections:
            risks = sections["risks_mitigation"]
            document.add_heading("Risks & Mitigation", 2)
            document.add_paragraph(risks)
            document.add_paragraph()

        if "resource_timeline" in sections:
            resources = sections["resource_timeline"]
            document.add_heading("Resource & Timeline", 2)
            document.add_paragraph(resources)
            document.add_paragraph()

        if "success_metrics" in sections:
            metrics = sections["success_metrics"]
            document.add_heading("Success Metrics", 2)
            document.add_paragraph(metrics)
            document.add_paragraph()

    document.add_page_break()
    document.add_heading("Approval & Sign-Off", 1)

    if doc.reviewed_by:
        document.add_paragraph(f"Reviewed by: {doc.reviewed_by}")
        document.add_paragraph(f"Review Status: {doc.status.upper()}")
        if doc.review_comment:
            document.add_paragraph(f"Comments: {doc.review_comment}")
    else:
        document.add_paragraph("Status: Pending Approval")

    docx_bytes = io.BytesIO()
    document.save(docx_bytes)
    docx_bytes.seek(0)

    return docx_bytes.getvalue()
